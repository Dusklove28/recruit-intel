"""Select a bounded, priority-ordered batch from the supplied history workbook."""

from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


@dataclass(frozen=True)
class HistoricalSeed:
    canonical_name: str
    historical_url: str
    candidate_root_url: str
    historical_platform: str
    suggested_domain: str
    historical_records: int
    batch_count: int
    seasons: str
    priority_tier: str


def _valid_web_url(value: object) -> str:
    url = str(value or "").strip()
    parsed = urlparse(url)
    return url if parsed.scheme in {"http", "https"} and parsed.hostname else ""


def load_priority_batch(path: Path, limit: int = 100) -> list[HistoricalSeed]:
    """Load only high-confidence historical seeds; historical URLs stay leads."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["高价值Seed"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        indices = {str(name): index for index, name in enumerate(headers)}
        required = {
            "公司名称", "历史记录数", "覆盖批次数", "出现批次", "历史代表URL",
            "候选根入口", "入口复用等级", "建议平台", "建议域名", "投递域名稳定度",
        }
        missing = required - indices.keys()
        if missing:
            raise ValueError(f"历史 Seed 表缺少字段：{', '.join(sorted(missing))}")

        selected: dict[str, HistoricalSeed] = {}
        for row in rows:
            get = lambda name: row[indices[name]] if indices[name] < len(row) else None
            name = str(get("公司名称") or "").strip()
            history_url = _valid_web_url(get("历史代表URL"))
            root_url = _valid_web_url(get("候选根入口"))
            tier = str(get("入口复用等级") or "")
            try:
                batches = int(get("覆盖批次数") or 0)
                records = int(get("历史记录数") or 0)
                stable_domain = float(get("投递域名稳定度") or 0) >= 1
            except (TypeError, ValueError):
                continue

            # This tab is the source's vetted high-value subset (P1 and >=2
            # seasons). Keep its A-tier, stable, non-WeChat entries first.
            candidate_urls = (history_url, root_url)
            if not name or batches < 2 or not tier.startswith("A") or not stable_domain:
                continue
            if not any(candidate_urls):
                continue
            if any("weixin.qq.com" in url.lower() or "mp.weixin" in url.lower() for url in candidate_urls):
                continue
            item = HistoricalSeed(
                canonical_name=name,
                historical_url=history_url,
                candidate_root_url=root_url,
                historical_platform=str(get("建议平台") or "unknown").strip(),
                suggested_domain=str(get("建议域名") or "").strip(),
                historical_records=records,
                batch_count=batches,
                seasons=str(get("出现批次") or "").strip(),
                priority_tier=tier,
            )
            existing = selected.get(name)
            if existing is None or (item.batch_count, item.historical_records) > (
                existing.batch_count, existing.historical_records
            ):
                selected[name] = item

        ranked = sorted(
            selected.values(),
            key=lambda item: (
                -item.batch_count,
                -item.historical_records,
                -bool(item.candidate_root_url),
                item.canonical_name,
            ),
        )
        if len(ranked) < limit:
            raise ValueError(f"符合优先条件的公司只有 {len(ranked)} 家，少于目标 {limit} 家")
        return ranked[:limit]
    finally:
        workbook.close()
