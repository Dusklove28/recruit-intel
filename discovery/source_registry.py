"""Persistent registry for currently verified employer recruitment sources."""

from dataclasses import asdict, dataclass
from datetime import date
import json
import sqlite3
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REGISTRY_FIELDS = (
    "canonical_name", "display_name", "aliases", "organization_type", "parent_group",
    "official_domain", "source_type", "current_source_url", "platform", "historical_url",
    "historical_platform", "discovered_by", "evidence_url", "verified_at", "status",
)
REGISTRY_STATUSES = frozenset({
    "verified", "expired_source", "redirected", "pending_manual_review",
    "access_control", "not_found", "unsupported",
})


@dataclass(frozen=True)
class SourceRecord:
    canonical_name: str
    display_name: str
    aliases: str = ""
    organization_type: str | None = None
    parent_group: str | None = None
    official_domain: str = ""
    source_type: str = ""
    current_source_url: str = ""
    platform: str = "unknown"
    historical_url: str = ""
    historical_platform: str = "unknown"
    discovered_by: str = "historical_seed"
    evidence_url: str = ""
    verified_at: str = ""
    status: str = "pending_manual_review"

    def validate(self) -> None:
        if self.status not in REGISTRY_STATUSES:
            raise ValueError(f"无效 Source Registry 状态：{self.status}")
        for field in ("current_source_url", "historical_url", "evidence_url"):
            value = getattr(self, field)
            if value and urlparse(value).scheme not in {"http", "https"}:
                raise ValueError(f"{field} 不是有效 HTTP(S) URL")


def ensure_source_registry(connection: sqlite3.Connection) -> None:
    connection.execute("""CREATE TABLE IF NOT EXISTS source_registry (
        canonical_name TEXT PRIMARY KEY,
        display_name TEXT NOT NULL,
        aliases TEXT NOT NULL DEFAULT '',
        organization_type TEXT,
        parent_group TEXT,
        official_domain TEXT NOT NULL DEFAULT '',
        source_type TEXT NOT NULL DEFAULT '',
        current_source_url TEXT NOT NULL DEFAULT '',
        platform TEXT NOT NULL DEFAULT 'unknown',
        historical_url TEXT NOT NULL DEFAULT '',
        historical_platform TEXT NOT NULL DEFAULT 'unknown',
        discovered_by TEXT NOT NULL DEFAULT '',
        evidence_url TEXT NOT NULL DEFAULT '',
        verified_at TEXT,
        status TEXT NOT NULL,
        updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    )""")
    connection.commit()


def upsert_source(connection: sqlite3.Connection, record: SourceRecord) -> None:
    record.validate()
    ensure_source_registry(connection)
    values = asdict(record)
    names = [*REGISTRY_FIELDS, "updated_at"]
    assignments = ", ".join(f"{name}=excluded.{name}" for name in names[1:])
    connection.execute(
        f"INSERT INTO source_registry ({', '.join(names)}) "
        f"VALUES ({', '.join('?' for _ in names[:-1])}, CURRENT_TIMESTAMP) "
        f"ON CONFLICT(canonical_name) DO UPDATE SET {assignments}",
        [values[name] for name in REGISTRY_FIELDS],
    )
    connection.commit()


def list_sources(connection: sqlite3.Connection) -> list[SourceRecord]:
    ensure_source_registry(connection)
    rows = connection.execute(
        f"SELECT {', '.join(REGISTRY_FIELDS)} FROM source_registry ORDER BY canonical_name"
    ).fetchall()
    return [SourceRecord(**dict(row)) for row in rows]


def export_source_registry(records: list[SourceRecord], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "招聘源核验"
    sheet.append(list(REGISTRY_FIELDS))
    url_fields = {"current_source_url", "historical_url", "evidence_url"}
    widths = [30, 30, 32, 16, 28, 28, 18, 52, 20, 52, 20, 20, 52, 16, 24]
    for item in records:
        item.validate()
        values = asdict(item)
        row = [values[name] for name in REGISTRY_FIELDS]
        sheet.append(row)
        for column, field in enumerate(REGISTRY_FIELDS, 1):
            cell = sheet.cell(sheet.max_row, column)
            if field in url_fields and cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(REGISTRY_FIELDS))}{sheet.max_row}"
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(path)
    return path


def dump_records(records: list[SourceRecord], path: Path) -> Path:
    """Write local runtime data; callers should place this under ignored data/."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps([asdict(item) for item in records], ensure_ascii=False, indent=2), encoding="utf-8")
    return path
