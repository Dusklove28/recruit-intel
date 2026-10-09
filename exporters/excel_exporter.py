"""A small Excel export with native clickable links and no extra columns."""

from pathlib import Path
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extractor.schema import FIELD_NAMES, RecruitmentRecord


WIDTHS = [8, 28, 15, 28, 24, 38, 27, 24, 48, 26, 15, 15, 14, 18, 18, 17, 45]
DATE_COLUMNS = {11, 12, 16}
LINK_COLUMNS = {14: "查看公告", 15: "立即报名"}


def export_excel(records: list[RecruitmentRecord], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "招聘汇总"
    sheet.append(FIELD_NAMES)

    for number, record in enumerate(records, start=1):
        values = record.model_dump(by_alias=True)
        values["序号"] = number
        sheet.append([values[name] for name in FIELD_NAMES])
        row_number = number + 1
        for column, label in LINK_COLUMNS.items():
            cell = sheet.cell(row_number, column)
            if cell.value:
                url = cell.value
                cell.value = label
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single")

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column in DATE_COLUMNS and cell.value is not None:
                cell.number_format = "yyyy-mm-dd"
        sheet.row_dimensions[row[0].row].height = 34
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:Q{max(2, sheet.max_row)}"

    temporary = path.with_name(path.stem + ".tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, path)
    return path
