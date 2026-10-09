"""Extract paragraphs and table rows from a DOCX attachment."""

from pathlib import Path

from docx import Document


def parse_docx(path: Path) -> str:
    document = Document(path)
    paragraphs = [item.text.strip() for item in document.paragraphs if item.text.strip()]
    tables = []
    for index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
        if rows:
            tables.append(f"表格 {index}\n" + "\n".join(rows))
    return "\n".join(paragraphs + tables)
