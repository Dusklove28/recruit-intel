"""Small offline checks for the new bounded discovery layer."""

import sqlite3
from openpyxl import load_workbook

from adapters.base import AccessRestricted, CampaignLead, DiscoveryOutcome, public_get
from adapters.beisen import BeisenAdapter
from adapters.generic import _cohort_evidence
from collectors.discovery import Candidate
from collectors.generic import Announcement
from config import Settings
from discovery.batch_runner import run_batch
from discovery.platform_detector import detect_platform
from discovery.seed_registry import Seed, STATUSES, load_seeds
from discovery.search_provider import MockSearchProvider, CandidateURL, rank_candidate, search_queries
from discovery.structured_records import beisen_draft, save_draft
from extractor.schema import RecruitmentRecord
from main import record_candidate_state
from processing.organization_registry import lookup_verified_child, seed_name_supported_by_notice, verify_child_seed
from storage.candidates import save_candidates, update_candidate
from storage.database import connect_database, save_record
from test_extraction import sample_payload


def _seed(name: str, url: str) -> Seed:
    return Seed(name, "央企", "国务院国资委", url, "example.cn", "custom", url, None)


def test_50_independent_seeds_and_platform_detection():
    seeds = load_seeds()
    assert len(seeds) == 50
    assert len({seed.career_url for seed in seeds}) == 50
    assert detect_platform("https://a.zhiye.com/campus/jobs") == "beisen"
    assert detect_platform("https://campus.51job.com/cofco/") == "51job_campus"
    assert detect_platform("https://job.example.cn/", official_domain="example.cn") == "custom"
    assert _cohort_evidence("中国华电2027年度校园招聘公告")
    assert _cohort_evidence("2027年度高校毕业生统招公告")
    assert not _cohort_evidence("2026年社会招聘公告")
    assert "access_control" in STATUSES


def test_search_fallback_is_only_an_interface():
    query = search_queries("中国大唐集团有限公司", "china-cdt.com")[0]
    result = CandidateURL("https://zhaopin.china-cdt.com/", "招聘", "mock", 100)
    assert MockSearchProvider({query: [result]}).search(query) == [result]
    assert rank_candidate(result.url, "china-cdt.com") > rank_candidate("https://example.com/notice", "china-cdt.com")


def test_curated_child_relation_requires_exact_official_chain(monkeypatch):
    seed = next(s for s in load_seeds() if s.organization_name == "中铁投资集团有限公司")
    monkeypatch.setattr("processing.organization_registry.central_enterprise_names", lambda: frozenset({"中国铁路工程集团有限公司"}))
    match = lookup_verified_child(seed)
    assert match is not None
    assert match.parent_unit == "中国铁路工程集团有限公司"
    assert len(match.parent_evidence) == 2
    assert lookup_verified_child(Seed("不相关公司", seed.organization_type, seed.parent_group,
                                      seed.career_url, seed.official_domain, seed.platform, seed.source, seed.verified_at)) is None


def test_abbreviated_name_needs_exact_legal_name_in_notice():
    official = "中国电力工程顾问集团中南电力设计院有限公司"
    assert seed_name_supported_by_notice("中南电力设计院有限公司", official, f"{official}2027届校园招聘")
    assert not seed_name_supported_by_notice("其他公司有限公司", official, f"{official}2027届校园招聘")
    assert not seed_name_supported_by_notice("中南电力设计院有限公司", official, "集团其他企业招聘")


def test_pending_candidate_retains_extracted_fields_and_evidence(tmp_path):
    settings = Settings("unused", "https://unused.example", "unused", tmp_path / "local.sqlite3",
                        tmp_path / "attachments", tmp_path / "out.xlsx")
    record = RecruitmentRecord.model_validate(sample_payload())
    notice = "https://official.example/notice/2027.html"
    record_candidate_state(settings, notice, notice, record, "待核验", "报名入口未证实",
                           {"学历要求": [notice]})
    with sqlite3.connect(settings.database_path) as connection:
        draft = connection.execute(
            "SELECT record_json,evidence_json FROM candidate_structured_records WHERE campaign_url=?", (notice,)
        ).fetchone()
        assert draft is not None
        assert '"学历要求"' in draft[1]
        assert '"education"' in draft[0]


def test_old_formal_row_is_not_counted_or_exported_after_pending_recheck(tmp_path, monkeypatch):
    notice = "https://jobs.example.cn/campus/2027.html"
    seed = _seed("测试集团", "https://jobs.example.cn/campus/index.html")
    settings = Settings("unused", "https://unused.example", "unused", tmp_path / "local.sqlite3",
                        tmp_path / "attachments", tmp_path / "out.xlsx")
    payload = sample_payload()
    payload.update({"单位类型": "央企", "招聘批次": "2027届校园招聘", "招聘对象": "2027届应届毕业生",
                    "官方公告": notice, "报名入口": "https://jobs.example.cn/apply", "招聘状态": "招聘中"})
    record = RecruitmentRecord.model_validate(payload)
    with connect_database(settings.database_path) as connection:
        save_record(connection, notice, record, "旧正文")
        save_candidates(connection, [Candidate("企业招聘入口", seed.career_url, "2027届校招", notice)])
        update_candidate(connection, notice, "待核验", reason="本次报名入口未证实")
    lead = CampaignLead(notice, "2027届校园招聘", seed.career_url, "2027届校园招聘")
    monkeypatch.setattr("discovery.batch_runner.GenericAdapter.discover",
                        lambda self, current: DiscoveryOutcome(current, "custom", True, [lead]))
    monkeypatch.setattr("discovery.batch_runner.public_get", lambda *args: object())
    monkeypatch.setattr("discovery.batch_runner.run", lambda *args, **kwargs: record)
    report, output = run_batch(settings, [seed], sleep_seconds=0)
    assert report.summary()["formal_campaigns"] == 0
    assert report.results[0].status == "pending_manual_review"
    assert load_workbook(output).active.max_row == 1


def test_access_control_is_recorded_and_next_seed_continues(tmp_path, monkeypatch):
    first = _seed("国家电网有限公司", "https://a.example.cn/campus")
    second = _seed("测试集团", "https://b.example.cn/campus")

    def discover(self, seed):
        if seed == first:
            raise AccessRestricted(seed.career_url, 412)
        return DiscoveryOutcome(seed, "custom", True, [], "未发现明确2027届校招活动")

    monkeypatch.setattr("discovery.batch_runner.GenericAdapter.discover", discover)
    settings = Settings("unused", "https://unused.example", "unused", tmp_path / "local.sqlite3", tmp_path / "attachments", tmp_path / "output" / "trial.xlsx")
    report, output = run_batch(settings, [first, second], sleep_seconds=0)
    assert output.exists()
    assert [item.status for item in report.results] == ["access_control", "no_2027_recruitment"]
    assert report.results[0].http_status == 412
    assert report.results[0].failed_url == first.career_url
    with sqlite3.connect(settings.database_path) as connection:
        rows = connection.execute("SELECT career_url,http_status,checked_at FROM career_seed_checks ORDER BY career_url").fetchall()
        assert len(rows) == 2
        assert rows[0][1] == 412
        assert rows[0][2]


def test_beisen_public_jobs_become_evidence_backed_draft():
    seed = _seed("测试集团", "https://demo.zhiye.com/campus/jobs")
    lead = CampaignLead(
        seed.career_url, "2027校园招聘", seed.career_url, "2027应届毕业生",
        structured_jobs=[{
            "JobAdName": "电气工程师", "Require": "2027届本科及以上应届毕业生，电气工程专业",
            "LocNames": ["北京"],
        }],
        listing_url="https://demo.zhiye.com/api/Jobad/GetJobAdPageList",
    )
    record, evidence = beisen_draft(seed, lead)
    assert record.education == "本科及以上"
    assert record.positions == "电气工程师"
    assert record.location == "北京"
    assert record.unit_type is None
    assert record.application_url is None
    assert evidence["education"] == [lead.listing_url]
    with sqlite3.connect(":memory:") as connection:
        save_draft(connection, lead, record, evidence)
        assert connection.execute("SELECT job_count FROM candidate_structured_records").fetchone()[0] == 1


def test_child_identity_requires_exact_official_relation(monkeypatch):
    source = "https://child.example.cn/campus/2027.html"
    seed = Seed("子公司甲有限公司", "央企子公司", "母集团有限公司", source, "example.cn", "custom", source, None)
    monkeypatch.setattr("processing.organization_registry.central_enterprise_names", lambda: frozenset({"母集团有限公司"}))
    monkeypatch.setattr(
        "processing.organization_registry.fetch_announcement",
        lambda url: Announcement(url, url, "子公司甲有限公司是母集团有限公司旗下全资子公司。", [], []),
    )
    match = verify_child_seed(seed)
    assert match is not None
    assert match.unit_type == "央企子公司"
    assert match.parent_evidence == (source,)
    monkeypatch.setattr(
        "processing.organization_registry.fetch_announcement",
        lambda url: Announcement(url, url, "子公司甲有限公司招聘。网站页脚：母集团有限公司", [], []),
    )
    assert verify_child_seed(seed) is None


def test_public_listing_with_login_modal_is_not_a_challenge():
    class Response:
        status_code = 200
        url = "https://example.cn/announcements"
        headers = {"Content-Type": "text/html"}
        encoding = "utf-8"
        text = "<title>招聘公告</title><main>2027校园招聘岗位</main><div>登录 验证码 captcha</div>"

        def raise_for_status(self):
            return None

    class Session:
        def get(self, *args, **kwargs):
            return Response()

    assert public_get(Session(), Response.url).status_code == 200


def test_beisen_redirect_to_another_tenant_is_not_followed(monkeypatch):
    seed = _seed("测试集团", "https://company.zhiye.com/campus/jobs")

    class Response:
        url = "https://other.zhiye.com/campus/jobs"
        text = ""

    monkeypatch.setattr("adapters.beisen.public_get", lambda session, url: Response())
    outcome = BeisenAdapter().discover(seed)
    assert outcome.leads == []
    assert "跳转" in outcome.reason
