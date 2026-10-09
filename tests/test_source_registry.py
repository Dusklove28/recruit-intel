import sqlite3

from openpyxl import load_workbook
import pytest

from discovery.source_registry import SourceRecord, export_source_registry, list_sources, upsert_source


def test_source_registry_persists_status_evidence_and_exports_clickable_urls(tmp_path):
    with sqlite3.connect(tmp_path / "registry.sqlite3") as connection:
        connection.row_factory = sqlite3.Row
        source = SourceRecord(
            canonical_name="甲集团有限公司",
            display_name="甲集团",
            official_domain="a.cn",
            source_type="企业招聘系统",
            current_source_url="https://job.a.cn/",
            historical_url="https://job.a.cn/campus/2026",
            historical_platform="beisen",
            platform="beisen",
            evidence_url="https://www.a.cn/about/recruitment.html",
            status="verified",
            verified_at="2026-10-09",
        )
        upsert_source(connection, source)
        assert list_sources(connection) == [source]

    path = export_source_registry([source], tmp_path / "source_registry.xlsx")
    sheet = load_workbook(path).active
    assert sheet.max_row == 2
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref.startswith("A1:O")
    assert sheet.cell(2, 8).hyperlink.target == "https://job.a.cn/"
    assert sheet.cell(2, 13).hyperlink.target == "https://www.a.cn/about/recruitment.html"


def test_source_record_rejects_unrecognized_status_and_invalid_url():
    with pytest.raises(ValueError):
        SourceRecord("甲", "甲", status="guessed").validate()
    with pytest.raises(ValueError):
        SourceRecord("甲", "甲", historical_url="javascript:alert(1)").validate()
