from datetime import date

from openpyxl import load_workbook

from exporters.excel_exporter import export_excel
from storage.database import connect_database, list_records, save_record

from test_extraction import sample_payload
from extractor.schema import RecruitmentRecord


def test_sqlite_upsert_and_excel_links(tmp_path):
    payload = sample_payload()
    payload["官方公告"] = "https://example.com/notice"
    record = RecruitmentRecord.model_validate(payload)
    database_path = tmp_path / "recruitment.sqlite3"
    with connect_database(database_path) as connection:
        save_record(connection, "https://example.com/notice", record, "原始正文")
        updated = record.model_copy(update={"notes": "复核完成", "verified_date": date(2026, 10, 9)})
        save_record(connection, "https://example.com/notice", updated, "新正文")
        records = list_records(connection)
        count = connection.execute("SELECT COUNT(*) FROM recruitment").fetchone()[0]
        raw_text = connection.execute("SELECT raw_text FROM recruitment").fetchone()[0]

    assert count == 1
    assert raw_text == "新正文"
    assert records[0].notes == "复核完成"
    assert records[0].updated_date == date(2026, 10, 9)

    output = export_excel(records, tmp_path / "out" / "招聘汇总.xlsx")
    sheet = load_workbook(output).active
    assert sheet.max_column == 17
    assert sheet.max_row == 2
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:Q2"
    assert sheet["N2"].hyperlink.target == "https://example.com/notice"
    assert sheet["O2"].hyperlink.target == "https://example.com/apply"
    assert sheet["P2"].number_format == "yyyy-mm-dd"


def test_content_date_changes_only_when_an_exported_content_field_changes(tmp_path):
    payload = sample_payload()
    payload["单位类型"] = None
    payload["最后核验日期"] = "2026-10-09"
    record = RecruitmentRecord.model_validate(payload)
    with connect_database(tmp_path / "recruitment.sqlite3") as connection:
        first = save_record(connection, "https://example.com/notice", record, "初始正文")
        assert first.updated_date == date(2026, 10, 9)
        verified_again = record.model_copy(update={"verified_date": date(2026, 10, 10)})
        second = save_record(connection, "https://example.com/notice", verified_again, "正文排版变化")
        assert second.updated_date == date(2026, 10, 9)
        assert second.verified_date == date(2026, 10, 10)
        changed = verified_again.model_copy(update={"unit_type": "央企", "parent_unit": "国务院国资委"})
        third = save_record(connection, "https://example.com/notice", changed, "正文排版变化")
        assert third.updated_date == date(2026, 10, 10)
        assert third.verified_date == date(2026, 10, 10)
